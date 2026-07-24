#!/usr/bin/env python3
"""Dump per-slide text from a .pptx so the skill can cite `slide N` anchors.

Self-contained, offline. Requires python-pptx (`pip install python-pptx`). If it's not
installed, this prints an install hint and exits non-zero — the skill then asks the user
to export the deck to PDF or paste key slides. It NEVER invents slide content.

Usage:
    python read_pptx.py deck.pptx            # print all slides
    python read_pptx.py deck.pptx 5          # print only slide 5
    python read_pptx.py deck.pptx 3-7        # print slides 3..7

Output format (stable, easy to anchor):
    ===== slide 5 =====
    <title / body text, notes>
"""
import sys


def _fail(msg, code=1):
    print(msg, file=sys.stderr)
    sys.exit(code)


def parse_range(arg, n):
    if "-" in arg:
        a, b = arg.split("-", 1)
        return range(int(a), int(b) + 1)
    k = int(arg)
    return range(k, k + 1)


def main():
    if len(sys.argv) < 2:
        _fail(__doc__)
    path = sys.argv[1]

    try:
        from pptx import Presentation
    except ImportError:
        _fail(
            "python-pptx not installed. Either:\n"
            "  pip install python-pptx\n"
            "or export the deck to PDF and read that instead, "
            "or paste the key slides. Do NOT guess slide content.",
            code=2,
        )

    try:
        prs = Presentation(path)
    except Exception as e:  # noqa: BLE001
        _fail(f"Could not open {path!r}: {e}")

    slides = list(prs.slides)
    total = len(slides)
    wanted = parse_range(sys.argv[2], total) if len(sys.argv) > 2 else range(1, total + 1)

    for i in wanted:
        if i < 1 or i > total:
            continue
        slide = slides[i - 1]
        print(f"===== slide {i} =====")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        print(line)
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if notes.strip():
            print(f"[notes] {notes.strip()}")
        print()


if __name__ == "__main__":
    main()
